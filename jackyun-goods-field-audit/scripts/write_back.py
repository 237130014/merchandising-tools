# -*- coding: utf-8 -*-
"""把抽查结果写回被核查表 —— 在表尾新增/复用「核查结果」「问题摘要」两列

核查结果取值：未通过（红）／通过（绿）／免检（灰，包材）
问题摘要：该行命中的所有疑点合并在一格，一个商品只看一行即可复核

用法:
    python write_back.py --src "10-15-2货品导出.xlsx" --issues "10-15-2_抽查疑点清单_20261016.xlsx"
    python write_back.py --src ... --issues ... --no-backup     # 跳过备份

注意：若被核查表正在 Excel 中打开，保存会报 PermissionError —— 先关闭 Excel。
"""
import argparse, collections, os, shutil, sys, datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULT_COL = "核查结果"
SUMMARY_COL = "问题摘要"
COL_CODE = 3     # 货品编号列（1-based），用于识别包材
COL_CLASS = 14   # 分类列（1-based）


def write_back(src, issues_path, sheet=None, backup=True, result_col=RESULT_COL, summary_col=SUMMARY_COL):
    if backup:
        bkdir = os.path.join(os.path.dirname(os.path.abspath(src)), ".workbuddy", "backup")
        os.makedirs(bkdir, exist_ok=True)
        bk = os.path.join(bkdir, f"{os.path.splitext(os.path.basename(src))[0]}_写回前备份_"
                                 f"{datetime.datetime.now():%Y%m%d%H%M}.xlsx")
        shutil.copy2(src, bk)
        print(f"[i] 已备份 → {bk}")

    wbi = openpyxl.load_workbook(issues_path, data_only=True)
    ws_i = wbi["疑点清单"] if "疑点清单" in wbi.sheetnames else wbi[wbi.sheetnames[0]]
    issues = list(ws_i.iter_rows(values_only=True))[1:]
    byrow = collections.OrderedDict()
    for r in issues:
        byrow.setdefault(r[0], []).append(f"[{r[4]}] {r[8]}")

    wb = openpyxl.load_workbook(src)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    ncol = ws.max_column
    if ncol >= 2 and str(ws.cell(1, ncol - 1).value or "").strip() == result_col \
       and str(ws.cell(1, ncol).value or "").strip() == summary_col:
        c_res, c_sum, action = ncol - 1, ncol, "复用已有列"
    else:
        c_res, c_sum, action = ncol + 1, ncol + 2, "新增列"
    print(f"[i] {action}：{get_column_letter(c_res)}({result_col}) / {get_column_letter(c_sum)}({summary_col})")

    HF = Font(bold=True, size=11)
    HFILL = PatternFill("solid", fgColor="FFF2CC")
    for c, t in ((c_res, result_col), (c_sum, summary_col)):
        cell = ws.cell(1, c, t)
        cell.font, cell.fill = HF, HFILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    F_RED, F_GREEN, F_GRAY = Font(color="C00000", bold=True), Font(color="1F7A3D"), Font(color="7F7F7F")
    FILL_RED = PatternFill("solid", fgColor="FCE4E4")
    FILL_GRAY = PatternFill("solid", fgColor="F2F2F2")
    FILL_NONE = PatternFill(fill_type=None)
    cnt = collections.Counter()

    for r in range(2, ws.max_row + 1):
        code = str(ws.cell(r, COL_CODE).value or "").strip()
        cls = str(ws.cell(r, COL_CLASS).value or "").strip()
        res, smy = ws.cell(r, c_res), ws.cell(r, c_sum)
        res.value, res.fill, res.font = None, FILL_NONE, Font()   # 清掉上一版结果
        smy.value = None
        if code[:2] == "BC" or cls == "包材":
            res.value, res.font, res.fill = "免检", F_GRAY, FILL_GRAY
            smy.value = "包材产品，跳过核对"
            cnt["免检"] += 1
        elif r in byrow:
            res.value, res.font, res.fill = "未通过", F_RED, FILL_RED
            smy.value = " ；".join(byrow[r])
            cnt["未通过"] += 1
        else:
            res.value, res.font = "通过", F_GREEN
            cnt["通过"] += 1
        res.alignment = Alignment(horizontal="center", vertical="center")
        smy.alignment = Alignment(vertical="center", wrap_text=True)

    ws.column_dimensions[get_column_letter(c_res)].width = 10
    ws.column_dimensions[get_column_letter(c_sum)].width = 62
    try:
        wb.save(src)
    except PermissionError:
        sys.exit(f"[ERROR] 写入被拒：{src}\n"
                 f"        该文件正被 Excel 打开，请先关闭 Excel 再重试。")
    print(f"[OK] 已写回 → {src}")
    print(f"     结果分布：{dict(cnt)}  合计 {sum(cnt.values())}")
    return cnt


def main():
    ap = argparse.ArgumentParser(description="把抽查结果写回被核查表")
    ap.add_argument("--src", required=True, help="被核查的货品导出表 (.xlsx)")
    ap.add_argument("--issues", required=True, help="疑点清单 (.xlsx)")
    ap.add_argument("--sheet", default=None, help="工作表名，默认第一个 sheet")
    ap.add_argument("--no-backup", action="store_true", help="不生成写回前备份")
    a = ap.parse_args()
    for p in (a.src, a.issues):
        if not os.path.isfile(p):
            sys.exit(f"[ERROR] 文件不存在: {p}")
    write_back(a.src, a.issues, a.sheet, backup=not a.no_backup)


if __name__ == "__main__":
    main()
